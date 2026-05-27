"""
generar_reporte.py
"""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
from pathlib import Path

# --- Configuración de Rutas de Producción ---
RAIZ_LOCAL = Path(__file__).parent.parent
FIGURAS_DIR = RAIZ_LOCAL / "data" / "figures"
PROCESADOS_DIR = RAIZ_LOCAL / "data" / "processed"
ARCHIVO_SALIDA_DOCX = RAIZ_LOCAL / "Reporte_Final_CNN_Splicing.docx"

class GeneradorDocumento:
    """
    Clase responsable de la creación del reporte técnico en formato DOCX.
    """

    def __init__(self):
        self.objeto_doc = Document()
        self._configurar_estilos_globales()

    def _configurar_estilos_globales(self):
        estilo_normal = self.objeto_doc.styles['Normal']
        estilo_normal.font.name = 'Calibri'
        estilo_normal.font.size = Pt(11)
        estilo_normal.paragraph_format.space_after = Pt(6)
        estilo_normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        for i in range(1, 4):
            estilo_h = self.objeto_doc.styles[f'Heading {i}']
            estilo_h.font.name = 'Calibri'
            estilo_h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            estilo_h.font.bold = True

    def agregar_titulo(self, texto_titulo: str, nivel_jerarquia: int = 1):
        self.objeto_doc.add_heading(texto_titulo, level=nivel_jerarquia)

    def agregar_parrafo(self, texto_parrafo: str):
        self.objeto_doc.add_paragraph(texto_parrafo)

    def agregar_imagen(self, ruta_imagen: Path, pie_foto: str = ""):
        if ruta_imagen.exists():
            self.objeto_doc.add_picture(str(ruta_imagen), width=Inches(5.5))
            if pie_foto:
                parrafo_pie = self.objeto_doc.add_paragraph()
                parrafo_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ejec_texto = parrafo_pie.add_run(pie_foto)
                ejec_texto.font.size = Pt(9)
                ejec_texto.font.italic = True
        else:
            self.agregar_parrafo(f"[Imagen {ruta_imagen.name} no encontrada]")

    def agregar_tabla(self, encabezados: list, filas: list):
        tabla = self.objeto_doc.add_table(rows=len(filas) + 1, cols=len(encabezados))
        tabla.style = 'Table Grid'
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(encabezados):
            celda = tabla.rows[0].cells[i]
            celda.text = h
            sombreado = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1F4E79" w:val="clear"/>')
            celda._tc.get_or_add_tcPr().append(sombreado)
        for i, fila in enumerate(filas):
            for j, val in enumerate(fila):
                tabla.rows[i+1].cells[j].text = str(val)

    def escribir_resultados(self):
        self.agregar_titulo("3. Resultados")
        self.agregar_parrafo("Visualizaciones detalladas del análisis genómico y desempeño del modelo:")
        
        figuras_lista = [
            ("gen_01_resumen_anotacion.png", "Composición global de la anotación genómica."),
            ("spec_01_balance_clases_splicing.png", "Balance de clases y tipos de sitio."),
            ("spec_04_contenido_gc_por_tipo.png", "Comparativa de Contenido GC (Donantes vs Aceptores)."),
            ("entrenamiento_curvas_desempeño.png", "Curvas de aprendizaje (Loss, Accuracy, F1)."),
            ("evaluacion_final_roc.png", "Curva ROC y AUC final.")
        ]
        
        for nombre_img, descripcion_img in figuras_lista:
            self.agregar_titulo(descripcion_img, nivel_jerarquia=2)
            self.agregar_imagen(FIGURAS_DIR / nombre_img, descripcion_img)

    def guardar_reporte(self):
        self.agregar_titulo("Reporte Final de Pipeline Genómico")
        self.escribir_resultados()
        self.objeto_doc.save(str(ARCHIVO_SALIDA_DOCX))
        print(f"✓ Reporte generado: {ARCHIVO_SALIDA_DOCX.name}")

if __name__ == "__main__":
    GeneradorDocumento().guardar_reporte()
